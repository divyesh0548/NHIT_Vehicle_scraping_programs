"""
Kerala registration date and fuel scraper.

Scrapes vehicle details from parivahan.gov.in Checkpost Tax portal for state
"KERALA" using the "VEHICLE TAX COLLECTION (OTHER STATE)" service.

Extracted fields from the Kerala tax collection page:
    1. Registration Date  -> input id "cal_regn_dt_input"
    2. Fuel               -> PrimeFaces selectOneMenu field labeled "Fuel"

Selenium mode:
    USE_SELENIUM_GRID defaults to config.SELENIUM_PROCESSING from .env.
    True  -> Selenium Grid / hub at SELENIUM_REMOTE_URL with parallel chunks.
            Parallelism = MAX_SELENIUM_GRID_NODES x SE_NODE_MAX_SESSIONS
            (e.g. 5 nodes x 5 sessions = 25 Chrome browsers).
    False -> single local Chrome browser.

Progress:
    Live-save mode updates the source workbook as each vehicle finishes, so a
    stopped run can resume from already-filled Registration Date / Fuel cells.
    Single-workbook mode can also keep the old per-node chunk_XX.xlsx progress
    files. Folder mode keeps a JSON status file in addition to live workbook
    updates.

Folder mode startup:
    1. Scan every workbook in the folder.
    2. Ensure each file has a JSON status entry (create if missing).
    3. Mark status "pending" when any eligible vehicle is missing data;
       mark "processed" when all eligible vehicles already have data.
    4. Process pending files one after another, scraping only missing rows
       without clearing existing Registration Date / Fuel values.
"""

import argparse
import json
import os
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from threading import Lock

# =============================================================================
# USER INPUTS - edit these variables and run the script directly.
# INPUT_PATH comes from REG_DATE_FUEL_EXCEL_PATH in .env (via config).
# =============================================================================
# Folder mode only. Relative names are created inside INPUT_PATH.
STATUS_FILE = "reg_date_fuel_status.json"

# Folder mode only. True scans subfolders too.
RECURSIVE = False

# Keep old per-node chunk progress files: "auto", "yes", or "no".
# auto = yes for a single workbook, no for folder mode.
PROGRESS_FILES = "auto"

# Update the source workbook while Selenium runs so stopping mid-process can resume.
LIVE_SAVE = True

# Save source workbook after this many processed rows. 1 is safest.
LIVE_SAVE_EVERY = max(1, int(os.getenv("REG_DATE_FUEL_LIVE_SAVE_EVERY", "1") or "1"))
# =============================================================================

import pandas as pd
from openpyxl import load_workbook
from selenium.common.exceptions import TimeoutException
from selenium.webdriver.common.by import By
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.ui import WebDriverWait

from config import (
    MAX_SELENIUM_GRID_NODES,
    REG_DATE_FUEL_EXCEL_PATH,
    SE_NODE_MAX_SESSIONS,
    SELENIUM_AUTO_MANAGE_NODES,
    SELENIUM_MAX_PARALLEL_SESSIONS,
    SELENIUM_PROCESSING,
    SELENIUM_REMOTE_URL,
)
from selenium_grid_manager import (
    assert_grid_ready,
    max_parallel_sessions,
    split_dataframe as split_df_for_grid,
    start_managed_nodes,
    stop_managed_nodes,
    wait_for_grid_ready,
)
from vehicle_number_utils import is_vehicle_number_eligible, normalize_vehicle_number

# Reuse the already-tested Kerala Checkpost navigation/browser helpers.
import web_scrape_kerala_checkpost as kerala_portal

# Can be a single .xlsx workbook or a folder containing .xlsx workbooks.
INPUT_PATH = REG_DATE_FUEL_EXCEL_PATH

WAIT_TIME = kerala_portal.WAIT_TIME
STATE_NAME = "KERALA"
SERVICE_NAME = "VEHICLE TAX COLLECTION (OTHER STATE)"
kerala_portal.STATE_NAME = STATE_NAME
kerala_portal.SERVICE_NAME = SERVICE_NAME

# .env switch:
#   SELENIUM_PROCESSING=true  -> Selenium Grid / hub
#   SELENIUM_PROCESSING=false -> normal local Chrome scraping
USE_SELENIUM_GRID = SELENIUM_PROCESSING

REGISTRATION_DATE_COLUMN = "Registration Date"
FUEL_COLUMN = "Fuel"
DETAIL_COLUMNS = [REGISTRATION_DATE_COLUMN, FUEL_COLUMN]
SUPPORTED_WORKBOOK_EXTENSIONS = {".xlsx"}
DEFAULT_STATUS_FILENAME = STATUS_FILE
DEFAULT_LIVE_SAVE_EVERY = LIVE_SAVE_EVERY

# Column used in per-node progress Excel files so merge can restore row positions.
ORIG_INDEX_COLUMN = "_orig_index"

EMPTY_DETAIL_VALUES = {
    "",
    "0",
    "-",
    "none",
    "nan",
    "n/a",
    "N/A",
    "Error",
    "Error - Input Not Found",
    "Error - Browser Restart Failed",
    "Error - Max Retries",
}


def _is_real_detail(value):
    """Return True when a scraped value looks like actual vehicle data."""
    if pd.isna(value):
        return False
    text = str(value).strip()
    if not text:
        return False
    lowered = text.lower()
    if lowered.startswith("---select"):
        return False
    return lowered not in {str(v).strip().lower() for v in EMPTY_DETAIL_VALUES}


def _is_completed_detail_value(value):
    """Return True when an output cell indicates this field was already attempted."""
    if pd.isna(value):
        return False
    text = str(value).strip()
    if not text:
        return False
    lowered = text.lower()
    if lowered.startswith("error"):
        return False
    return True


def _row_has_completed_details(row):
    return all(
        col in row and _is_completed_detail_value(row.get(col, ""))
        for col in DETAIL_COLUMNS
    )


def _set_details(df, idx, details):
    for col, value in details.items():
        if col not in df.columns:
            df[col] = ""
        df.at[idx, col] = value


def _na_details():
    return {col: "N/A" for col in DETAIL_COLUMNS}


def _error_details(message):
    return {col: message for col in DETAIL_COLUMNS}


def _read_input_by_label(driver, label_text):
    """Read the input value from the grid cell that has the given label text."""
    try:
        value = driver.execute_script(
            """
            var labelText = arguments[0];
            var root = document.getElementById('kltaxcollection') || document;

            function norm(s) {
              return (s || '').replace(/\\s+/g, ' ').trim();
            }

            var labels = root.querySelectorAll('span.ui-outputlabel-label');
            for (var i = 0; i < labels.length; i++) {
              if (norm(labels[i].textContent) !== labelText) continue;
              var col = labels[i].closest('[class*="ui-grid-col"]');
              if (!col) continue;
              var input = col.querySelector('input');
              if (!input) continue;
              return norm(input.value || '');
            }
            return '';
            """,
            label_text,
        )
        return (value or "").strip()
    except Exception:
        return ""


def _read_registration_date(driver, quiet=False):
    """Read Registration Date from the Kerala details panel."""
    for field_id in ("cal_regn_dt_input",):
        try:
            element = driver.find_element(By.ID, field_id)
            value = (element.get_attribute("value") or "").strip()
            if _is_real_detail(value):
                if not quiet:
                    print(f"  {field_id}: {value}")
                return value
        except Exception:
            continue

    value = _read_input_by_label(driver, "Registration Date")
    if _is_real_detail(value):
        if not quiet:
            print(f"  Registration Date via label: {value}")
        return value

    return ""


def _selected_option_text(driver, select_or_label_id):
    """Read visible/selected text from a PrimeFaces select or label id."""
    try:
        element = driver.find_element(By.ID, select_or_label_id)
        tag = (element.tag_name or "").lower()
        if tag == "select":
            text = driver.execute_script(
                """
                var s = arguments[0];
                if (!s || s.selectedIndex < 0) return '';
                var opt = s.options[s.selectedIndex];
                return opt ? (opt.textContent || opt.text || '').trim() : '';
                """,
                element,
            )
        else:
            text = driver.execute_script(
                "return (arguments[0].textContent || arguments[0].innerText || '').trim();",
                element,
            )
        return (text or "").strip()
    except Exception:
        return ""


def _read_fuel(driver, quiet=False):
    """Read Fuel from the Kerala PrimeFaces dropdown."""
    value = kerala_portal._js_read_dropdown_by_label(driver, "Fuel", quiet=True)
    if _is_real_detail(value):
        if not quiet:
            print(f"  Fuel: {value} (via label JS)")
        return value

    # The reference page uses j_idt342, but these j_idt* ids can move between
    # deployments. Treat them as useful fallbacks only.
    for element_id in ("j_idt342_input", "j_idt342_label"):
        value = _selected_option_text(driver, element_id)
        if _is_real_detail(value):
            if not quiet:
                print(f"  {element_id}: {value}")
            return value

    if not quiet:
        value = kerala_portal._read_select_label(
            driver,
            ["j_idt342_label", "j_idt342_input"],
            "Fuel",
        )
        if _is_real_detail(value):
            return value

    return ""


def _vehicle_details_ready(driver):
    """
    True when at least one target vehicle-specific field has loaded.

    This avoids trusting pre-filled page-level fields such as From State.
    """
    return bool(
        _read_registration_date(driver, quiet=True)
        or _read_fuel(driver, quiet=True)
    )


def _wait_for_registration_and_fuel(driver, timeout=12):
    """Wait for target fields to populate after Get Details."""
    deadline = time.time() + timeout
    last_registration_date = ""
    last_fuel = ""

    print("  Waiting for Registration Date / Fuel to populate...")
    while time.time() < deadline:
        last_registration_date = _read_registration_date(driver, quiet=True)
        last_fuel = _read_fuel(driver, quiet=True)
        if last_registration_date and last_fuel:
            print(f"  Registration Date: {last_registration_date}")
            print(f"  Fuel: {last_fuel}")
            return last_registration_date, last_fuel
        time.sleep(0.5)

    if last_registration_date or last_fuel:
        print(
            "  [WARN] Timed out waiting for both fields "
            f"(registration='{last_registration_date or 'none'}', fuel='{last_fuel or 'none'}')"
        )
    else:
        print("  [WARN] Timed out waiting for Registration Date / Fuel")

    return last_registration_date, last_fuel


def _details_look_populated(details):
    """True if scraped details look like real vehicle data."""
    return any(_is_real_detail(details.get(col, "")) for col in DETAIL_COLUMNS)


def get_vehicle_details(driver):
    """Extract Registration Date and Fuel from the loaded Kerala details page."""
    print("Getting Kerala registration/fuel details...")

    registration_date, fuel = _wait_for_registration_and_fuel(driver, timeout=12)

    if not registration_date:
        registration_date = _read_registration_date(driver)
    if not fuel:
        fuel = _read_fuel(driver)

    details = {
        REGISTRATION_DATE_COLUMN: registration_date if registration_date else "none",
        FUEL_COLUMN: fuel if fuel else "none",
    }
    print(
        f"Registration Date: {details[REGISTRATION_DATE_COLUMN]} | "
        f"Fuel: {details[FUEL_COLUMN]}"
    )
    return details


def restart_browser_and_continue(driver):
    """Restart browser and re-navigate to the Kerala tax page."""
    print("\n[RESTART] RESTARTING BROWSER...")
    try:
        driver.quit()
        print("Closed current browser session")
    except Exception:
        print("Could not properly close browser")

    time.sleep(3)

    new_driver = kerala_portal.setup_driver()
    if not new_driver:
        print("Could not restart browser")
        return None, None

    new_wait = WebDriverWait(new_driver, WAIT_TIME)

    print("Re-navigating to tax page...")
    if not kerala_portal.navigate_to_tax_page(new_driver, new_wait):
        print("Could not navigate to tax page after restart")
        try:
            new_driver.quit()
        except Exception:
            pass
        return None, None

    print("[OK] Browser restarted successfully")
    return new_driver, new_wait


def process_single_vehicle(driver, wait, vehicle_no, idx, df):
    """Process one vehicle and extract Registration Date + Fuel."""
    max_retries = 5
    retry_count = 0

    while retry_count <= max_retries:
        try:
            print(f"\n{'=' * 50}")
            print(
                f"Processing row {idx} of {len(df)} - Vehicle: {vehicle_no} "
                f"(Attempt {retry_count + 1})"
            )
            print(f"{'=' * 50}")

            print("Quick refreshing page...")
            driver.execute_script("location.reload()")
            time.sleep(2)

            vehicle_input_xpath = (
                "//input[@type='text' and @maxlength='10' and not(@id='mobileno')]"
            )

            print("Waiting for Vehicle Number input to be interactable...")
            try:
                input_element = WebDriverWait(driver, 15).until(
                    EC.element_to_be_clickable((By.XPATH, vehicle_input_xpath))
                )
                print("[OK] Vehicle Number input found and ready")
            except TimeoutException:
                print("[ERROR] Timeout: could not find Vehicle Number input")
                if retry_count < max_retries:
                    retry_count += 1
                    print(
                        f"[RETRY] Browser restart "
                        f"({retry_count}/{max_retries})..."
                    )
                    new_driver, new_wait = restart_browser_and_continue(driver)
                    if new_driver:
                        driver = new_driver
                        wait = new_wait
                        continue
                    break
                _set_details(df, idx, _error_details("Error - Input Not Found"))
                return driver, wait, False

            try:
                popup_xpath = (
                    "//div[contains(@class,'ui-dialog') and contains(@style,'display: block')]"
                )
                popup = WebDriverWait(driver, 2).until(
                    EC.visibility_of_element_located((By.XPATH, popup_xpath))
                )
                if popup:
                    ok_button_xpath = (
                        "//div[contains(@class,'ui-dialog')]"
                        "//button[.//span[contains(text(),'OK') or contains(text(),'Ok')]]"
                    )
                    kerala_portal.safe_click(driver, wait, ok_button_xpath, "OK button on popup")
            except TimeoutException:
                pass

            try:
                input_element = WebDriverWait(driver, 5).until(
                    EC.element_to_be_clickable((By.XPATH, vehicle_input_xpath))
                )
                driver.execute_script(
                    """
                    var el = arguments[0], val = arguments[1];
                    el.value = val;
                    el.dispatchEvent(new Event('input', {bubbles: true}));
                    el.dispatchEvent(new Event('change', {bubbles: true}));
                    if (typeof el.onkeyup === 'function') {
                      try { el.onkeyup(); } catch (e) {}
                    }
                    """,
                    input_element,
                    vehicle_no,
                )
                print("[OK] Vehicle number entered via JavaScript")
            except Exception:
                try:
                    input_element.clear()
                    input_element.send_keys(vehicle_no)
                    print("[OK] Vehicle number entered")
                except Exception as e:
                    print(f"[ERROR] Could not enter vehicle number: {e}")
                    retry_count += 1
                    continue

            get_details_xpath = "//button[.//span[contains(text(), 'Get Details')]]"
            try:
                get_details_btn = WebDriverWait(driver, 5).until(
                    EC.element_to_be_clickable((By.XPATH, get_details_xpath))
                )
                driver.execute_script("arguments[0].click();", get_details_btn)
                print("[OK] Get Details clicked")
            except Exception:
                kerala_portal.safe_click(driver, wait, get_details_xpath, "Get Details button")

            print("Waiting for vehicle details...")
            time.sleep(3)

            popup_appeared = False
            data_appeared = False

            for _ in range(8):
                try:
                    popup_xpath = (
                        "//div[contains(@class,'ui-dialog') and "
                        "contains(@style,'display: block')]"
                    )
                    try:
                        WebDriverWait(driver, 2).until(
                            EC.visibility_of_element_located((By.XPATH, popup_xpath))
                        )
                        popup_appeared = True
                        print("Popup detected - dismissing and trying to read details")
                        ok_button_xpath = (
                            "//div[contains(@class,'ui-dialog')]"
                            "//button[.//span[contains(text(),'OK') or contains(text(),'Ok')]]"
                        )
                        kerala_portal.safe_click(
                            driver, wait, ok_button_xpath, "OK button on popup"
                        )
                        time.sleep(1)
                    except TimeoutException:
                        pass

                    if _vehicle_details_ready(driver):
                        data_appeared = True
                        print("[OK] Vehicle details loaded")
                        break
                except Exception:
                    pass

                time.sleep(1)

            if data_appeared:
                details = get_vehicle_details(driver)
                if _details_look_populated(details):
                    _set_details(df, idx, details)
                    print(f"[OK] Details extracted for {vehicle_no}")
                    return driver, wait, True
                print("[WARNING] Details signal present but extracted values were empty")

            if popup_appeared and not data_appeared:
                _set_details(df, idx, _na_details())
                print("[OK] Marked as N/A (popup only, no usable details)")
                return driver, wait, True

            if retry_count < max_retries:
                retry_count += 1
                print(
                    f"[WARNING] No usable details for {vehicle_no} - "
                    f"retrying on same browser ({retry_count}/{max_retries})..."
                )
                continue

            _set_details(df, idx, _na_details())
            print(
                f"[OK] Marked as N/A for {vehicle_no} "
                f"(no details after {max_retries + 1} attempts)"
            )
            return driver, wait, True

        except Exception as e:
            print(f"Error processing {vehicle_no}: {e}")
            if retry_count < max_retries:
                retry_count += 1
                print(f"[RETRY] Browser restart ({retry_count}/{max_retries})...")
                new_driver, new_wait = restart_browser_and_continue(driver)
                if new_driver:
                    driver = new_driver
                    wait = new_wait
                    continue
                break
            _set_details(df, idx, _error_details("Error"))
            return driver, wait, False

    _set_details(df, idx, _error_details("Error - Max Retries"))
    return driver, wait, False


def make_progress_dir(base_path=None):
    """Create a directory for per-node progress Excel files."""
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    if base_path:
        base = Path(base_path)
        parent = base.parent if base.suffix else base
        name = (
            f"{base.stem}_reg_date_fuel_progress_{stamp}"
            if base.suffix
            else f"reg_date_fuel_progress_{stamp}"
        )
        progress_dir = parent / name
    else:
        progress_dir = Path(f"kerala_reg_date_fuel_progress_{stamp}")
    progress_dir.mkdir(parents=True, exist_ok=True)
    print(f"[PROGRESS] Writing outputs to: {progress_dir}")
    return str(progress_dir)


def chunk_progress_path(progress_dir, chunk_id):
    """Path for one Selenium Grid node's progress Excel file."""
    return str(Path(progress_dir) / f"chunk_{int(chunk_id):02d}.xlsx")


class ExcelLiveUpdater:
    """Thread-safe, atomic row updater for the source workbook."""

    def __init__(self, excel_path, sheet_name, detail_columns=None, save_every=1):
        self.excel_path = Path(excel_path)
        self.sheet_name = sheet_name
        self.detail_columns = list(detail_columns or DETAIL_COLUMNS)
        self.save_every = max(1, int(save_every or 1))
        self.lock = Lock()
        self.pending_saves = 0
        self.workbook = load_workbook(self.excel_path)
        if sheet_name not in self.workbook.sheetnames:
            raise ValueError(f"Sheet not found in workbook: {sheet_name}")
        self.sheet = self.workbook[sheet_name]
        self.column_indexes = self._ensure_detail_headers()
        self._save_locked(reason="headers")
        print(
            f"[LIVE SAVE] Updating source workbook after every "
            f"{self.save_every} processed row(s): {self.excel_path} [{sheet_name}]",
            flush=True,
        )

    def _ensure_detail_headers(self):
        headers = {}
        for cell in self.sheet[1]:
            value = str(cell.value or "").strip()
            if value:
                headers[value] = cell.column

        next_col = self.sheet.max_column + 1
        indexes = {}
        for column_name in self.detail_columns:
            if column_name not in headers:
                self.sheet.cell(row=1, column=next_col).value = column_name
                headers[column_name] = next_col
                next_col += 1
            indexes[column_name] = headers[column_name]
        return indexes

    def _save_locked(self, reason="progress"):
        tmp_path = self.excel_path.with_name(
            f".{self.excel_path.stem}.live_save.tmp{self.excel_path.suffix}"
        )
        self.workbook.save(tmp_path)
        tmp_path.replace(self.excel_path)
        self.pending_saves = 0
        print(f"[LIVE SAVE] Saved {reason}: {self.excel_path}", flush=True)

    def update_row(self, row_index, details):
        """Update only the scraped row's detail columns; never wipe other rows."""
        with self.lock:
            excel_row = int(row_index) + 2
            for column_name in self.detail_columns:
                if column_name not in self.column_indexes:
                    continue
                new_value = details.get(column_name, "")
                if new_value is None:
                    new_value = ""
                cell = self.sheet.cell(
                    row=excel_row,
                    column=self.column_indexes[column_name],
                )
                # Do not blank an already-filled cell with an empty update.
                existing = "" if cell.value is None else str(cell.value).strip()
                if existing and not str(new_value).strip():
                    continue
                cell.value = new_value
            self.pending_saves += 1
            if self.pending_saves >= self.save_every:
                self._save_locked(reason=f"row {excel_row}")

    def flush(self):
        with self.lock:
            if self.pending_saves:
                self._save_locked(reason="final flush")


def save_chunk_progress(df, progress_path):
    """Persist current chunk DataFrame to its own Excel file."""
    if not progress_path:
        return
    try:
        out = df.copy()
        out[ORIG_INDEX_COLUMN] = out.index
        cols = [ORIG_INDEX_COLUMN] + [c for c in out.columns if c != ORIG_INDEX_COLUMN]
        out = out[cols]
        Path(progress_path).parent.mkdir(parents=True, exist_ok=True)
        out.to_excel(progress_path, index=False)
    except Exception as e:
        print(f"  [PROGRESS] Failed to save {progress_path}: {e}", flush=True)


def load_chunk_progress_frame(progress_path):
    """Load a chunk progress Excel and restore original index."""
    frame = pd.read_excel(progress_path, dtype=str)
    if ORIG_INDEX_COLUMN in frame.columns:
        idx = pd.to_numeric(frame[ORIG_INDEX_COLUMN], errors="coerce")
        frame = frame.drop(columns=[ORIG_INDEX_COLUMN])
        valid = idx.notna()
        frame = frame.loc[valid].copy()
        frame.index = idx.loc[valid].astype(int)
    return frame


def merge_chunk_progress_files(df_input, progress_dir):
    """Merge all chunk_*.xlsx files from progress_dir onto a copy of df_input."""
    if not progress_dir:
        return None

    progress_path = Path(progress_dir)
    if not progress_path.exists():
        return None

    chunk_files = sorted(progress_path.glob("chunk_*.xlsx"))
    if not chunk_files:
        print(f"[PROGRESS] No chunk files found in {progress_dir}")
        return None

    merged = df_input.copy()
    for col in DETAIL_COLUMNS:
        if col not in merged.columns:
            merged[col] = ""

    loaded = 0
    for path in chunk_files:
        try:
            frame = load_chunk_progress_frame(path)
            common = frame.index.intersection(merged.index)
            if len(common) == 0:
                print(f"[PROGRESS] Warning: no index overlap for {path.name}; skipped")
                continue
            for col in frame.columns:
                if col not in merged.columns:
                    merged[col] = ""
                merged.loc[common, col] = frame.loc[common, col]
            loaded += 1
            print(f"[PROGRESS] Merged {path.name} ({len(common)} rows)")
        except Exception as e:
            print(f"[PROGRESS] Could not merge {path}: {e}")

    if loaded == 0:
        return None

    merged_out = progress_path / "merged_output.xlsx"
    try:
        merged.to_excel(merged_out, index=False)
        print(f"[PROGRESS] Wrote merged output: {merged_out}")
    except Exception as e:
        print(f"[PROGRESS] Could not write merged_output.xlsx: {e}")

    return merged


def find_vehicle_reg_column(df):
    """
    Return the vehicle registration column if a matching header keyword is found.

    Prefers columns containing both 'veh' and 'reg', then falls back to any
    column containing 'vehicle'.
    """
    for col in df.columns:
        lower = str(col).lower()
        if "veh" in lower and "reg" in lower:
            return col
    for col in df.columns:
        if "vehicle" in str(col).lower():
            return col
    return None


def _scrape_kerala_reg_date_fuel_impl(
    df_input,
    progress_path=None,
    progress_callback=None,
):
    df = df_input.copy()

    veh_col = find_vehicle_reg_column(df)
    if veh_col is None:
        print("[ERROR] Could not find vehicle number column in dataframe")
        save_chunk_progress(df, progress_path)
        return df

    for col in DETAIL_COLUMNS:
        if col not in df.columns:
            df[col] = ""

    df[veh_col] = df[veh_col].apply(normalize_vehicle_number)
    invalid_vehicle_count = int((~df[veh_col].apply(is_vehicle_number_eligible)).sum())
    if invalid_vehicle_count:
        print(
            f"[INFO] Skipping {invalid_vehicle_count} vehicle numbers outside eligible "
            f"length range (allowed lengths: 8, 9, 10)"
        )

    print(f"Total vehicles to process: {len(df)}")
    if len(df) == 0:
        print("No vehicles to process.")
        save_chunk_progress(df, progress_path)
        return df

    valid_mask = df[veh_col].apply(is_vehicle_number_eligible)
    completed_mask = df.apply(_row_has_completed_details, axis=1)
    completed_count = int((valid_mask & completed_mask).sum())
    if completed_count:
        print(
            f"[RESUME] Skipping {completed_count} vehicle(s) already filled in "
            f"{REGISTRATION_DATE_COLUMN}/{FUEL_COLUMN}"
        )
    remaining_df = df[valid_mask & ~completed_mask]
    remaining_count = len(remaining_df)

    print("\n" + "=" * 80)
    print(f"Web scraping {remaining_count} vehicles for Registration Date and Fuel...")
    print("=" * 80)

    save_chunk_progress(df, progress_path)

    if remaining_count == 0:
        print("No valid vehicles to scrape.")
        return df

    driver = kerala_portal.setup_driver()
    if not driver:
        print("[ERROR] Could not start browser - aborting")
        save_chunk_progress(df, progress_path)
        return df

    wait = WebDriverWait(driver, WAIT_TIME)

    print("Navigating to Kerala tax page...")
    max_nav_attempts = 5
    nav_success = False

    for nav_attempt in range(1, max_nav_attempts + 1):
        print(f"[Attempt {nav_attempt}/{max_nav_attempts}] Navigating to tax page...")
        if kerala_portal.navigate_to_tax_page(driver, wait):
            nav_success = True
            print("[OK] Navigation successful")
            break

        if nav_attempt < max_nav_attempts:
            print("Restarting browser and retrying...")
            try:
                driver.quit()
            except Exception:
                pass
            time.sleep(5)
            driver = kerala_portal.setup_driver()
            if not driver:
                break
            wait = WebDriverWait(driver, WAIT_TIME)

    if not nav_success:
        print("[ERROR] Could not navigate to tax page - aborting")
        try:
            driver.quit()
        except Exception:
            pass
        save_chunk_progress(df, progress_path)
        return df

    start_time = time.perf_counter()
    scraped_count = 0

    for idx, row in remaining_df.iterrows():
        vehicle_no = normalize_vehicle_number(row[veh_col])
        if not vehicle_no or not is_vehicle_number_eligible(vehicle_no):
            continue

        driver, wait, success = process_single_vehicle(driver, wait, vehicle_no, idx, df)

        save_chunk_progress(df, progress_path)
        if progress_callback:
            progress_callback(idx, {col: df.at[idx, col] for col in DETAIL_COLUMNS})

        if success:
            scraped_count += 1

        print(f"Progress: {scraped_count}/{remaining_count} scraped")
        time.sleep(1)

    try:
        driver.quit()
    except Exception:
        pass

    save_chunk_progress(df, progress_path)

    elapsed = time.perf_counter() - start_time
    hrs = int(elapsed // 3600)
    mins = int((elapsed % 3600) // 60)
    secs = int(elapsed % 60)
    print(f"\nWeb scraping completed in {hrs:02d}:{mins:02d}:{secs:02d}")

    print("\n" + "=" * 80)
    print("PROCESSING COMPLETE")
    print("=" * 80)
    print(f"Total vehicles processed: {len(df)}")
    print(f"Web scraped: {scraped_count}")
    if progress_path:
        print(f"Chunk progress file: {progress_path}")
    print("=" * 80)

    return df


def _scrape_chunk(
    chunk_df,
    remote_url,
    chunk_id=None,
    progress_path=None,
    progress_callback=None,
):
    """Run scraping for one chunk with a per-thread remote URL."""
    kerala_portal._thread_remote_url.value = remote_url
    try:
        if progress_path:
            print(
                f"[PROGRESS] Node/chunk {chunk_id} output file: {progress_path}",
                flush=True,
            )
        return _scrape_kerala_reg_date_fuel_impl(
            chunk_df,
            progress_path=progress_path,
            progress_callback=progress_callback,
        )
    finally:
        kerala_portal._thread_remote_url.value = None


def _run_grid_scrape(
    df_input,
    remote_url,
    progress_dir=None,
    write_progress_files=True,
    progress_callback=None,
):
    """
    Split work across Selenium Grid sessions and merge per-chunk progress files.

    Layout from .env:
      MAX_SELENIUM_GRID_NODES chrome node containers
      SE_NODE_MAX_SESSIONS Chrome browsers per node
      => SELENIUM_MAX_PARALLEL_SESSIONS total parallel scrapers
    """
    parallel_sessions = max_parallel_sessions()
    chunks = split_df_for_grid(df_input, parallel_sessions)
    if not chunks:
        return df_input

    if write_progress_files and progress_dir is None:
        progress_dir = make_progress_dir()
    elif write_progress_files:
        Path(progress_dir).mkdir(parents=True, exist_ok=True)

    print(
        f"  [SELENIUM GRID] {len(chunks)} session chunk(s) on "
        f"{MAX_SELENIUM_GRID_NODES} node(s) x {SE_NODE_MAX_SESSIONS} session(s) "
        f"(cap {SELENIUM_MAX_PARALLEL_SESSIONS}) -> {remote_url} "
        f"| progress_dir={progress_dir if write_progress_files else 'disabled'}",
        flush=True,
    )

    managed_nodes = []
    merged = None
    grid_error = None

    try:
        if SELENIUM_AUTO_MANAGE_NODES:
            # Start node containers only (not one container per session).
            managed_nodes = start_managed_nodes(MAX_SELENIUM_GRID_NODES)
            wait_for_grid_ready(remote_url)
        else:
            assert_grid_ready(remote_url)

        result_frames = []
        failures = []
        with ThreadPoolExecutor(max_workers=len(chunks)) as executor:
            future_to_chunk = {}
            for chunk_id, chunk in enumerate(chunks, start=1):
                path = (
                    chunk_progress_path(progress_dir, chunk_id)
                    if write_progress_files
                    else None
                )
                if path:
                    save_chunk_progress(chunk, path)
                future = executor.submit(
                    _scrape_chunk,
                    chunk,
                    remote_url,
                    chunk_id,
                    path,
                    progress_callback,
                )
                future_to_chunk[future] = chunk_id

            for future in as_completed(future_to_chunk):
                chunk_id = future_to_chunk[future]
                try:
                    result_frames.append(future.result())
                    print(
                        f"  [SELENIUM GRID] chunk {chunk_id}/{len(chunks)} completed",
                        flush=True,
                    )
                except Exception as exc:
                    failures.append((chunk_id, exc))
                    print(
                        f"  [SELENIUM GRID] chunk {chunk_id} failed: {exc}",
                        flush=True,
                    )

        merged = (
            merge_chunk_progress_files(df_input, progress_dir)
            if write_progress_files
            else None
        )
        if merged is None and result_frames:
            merged = df_input.copy()
            for col in DETAIL_COLUMNS:
                if col not in merged.columns:
                    merged[col] = ""
            for frame in result_frames:
                for col in frame.columns:
                    merged.loc[frame.index, col] = frame[col]

        if not result_frames and merged is None:
            grid_error = f"All Selenium Grid chunks failed: {failures}"
        elif failures and merged is None:
            grid_error = f"{len(failures)} Selenium Grid chunk(s) failed: {failures}"
        elif failures:
            print(
                f"  [SELENIUM GRID] {len(failures)} chunk(s) had errors; "
                f"keeping merged progress from disk",
                flush=True,
            )

    except Exception as exc:
        grid_error = str(exc)
        print(f"  [SELENIUM GRID] error: {exc}", flush=True)
        recovered = (
            merge_chunk_progress_files(df_input, progress_dir)
            if write_progress_files
            else None
        )
        if recovered is not None:
            merged = recovered
    finally:
        if managed_nodes:
            stop_managed_nodes(managed_nodes)

    if grid_error and merged is None:
        print(
            f"  [SELENIUM GRID] falling back to local Chrome... ({grid_error})",
            flush=True,
        )
        local_path = (
            chunk_progress_path(progress_dir, 99)
            if write_progress_files
            else None
        )
        return _scrape_chunk(
            df_input,
            None,
            chunk_id=99,
            progress_path=local_path,
            progress_callback=progress_callback,
        )

    if grid_error and merged is not None:
        if write_progress_files:
            print(
                f"  [SELENIUM GRID] partial progress kept from {progress_dir}",
                flush=True,
            )
        else:
            print("  [SELENIUM GRID] partial in-memory progress kept", flush=True)

    return merged if merged is not None else df_input


def scrape_kerala_reg_date_fuel(
    df_input,
    remote_url=None,
    use_selenium_grid=None,
    progress_dir=None,
    write_progress_files=True,
    progress_callback=None,
):
    """
    Scrape Registration Date and Fuel for all eligible vehicles in df_input.

    use_selenium_grid:
        True  -> Selenium Grid at SELENIUM_REMOTE_URL
        False -> local Chrome
        None  -> module-level USE_SELENIUM_GRID, sourced from SELENIUM_PROCESSING
    remote_url:
        Explicit Grid URL. When set, it forces grid usage.
    progress_dir:
        Folder for per-node chunk_XX.xlsx progress files.
    write_progress_files:
        True keeps the old chunk_XX.xlsx progress files. False keeps progress
        in memory for callers that update each source workbook directly.
    """
    if use_selenium_grid is None:
        use_selenium_grid = USE_SELENIUM_GRID

    use_grid = use_selenium_grid or (remote_url is not None)
    grid_url = remote_url or SELENIUM_REMOTE_URL

    print("=" * 80)
    print("STARTING KERALA REGISTRATION DATE / FUEL SCRAPING")
    print("=" * 80)
    print(f"State: {STATE_NAME}")
    print(f"Service: {SERVICE_NAME}")

    if not use_grid:
        print("Web scrape mode: local Chrome (single browser)")
        if write_progress_files and progress_dir is None:
            progress_dir = make_progress_dir()
        local_path = (
            chunk_progress_path(progress_dir, 1)
            if write_progress_files
            else None
        )
        return _scrape_chunk(
            df_input,
            None,
            chunk_id=1,
            progress_path=local_path,
            progress_callback=progress_callback,
        )

    if write_progress_files and progress_dir is None:
        progress_dir = make_progress_dir()
    elif write_progress_files:
        Path(progress_dir).mkdir(parents=True, exist_ok=True)
        print(f"[PROGRESS] Writing outputs to: {progress_dir}")
    else:
        print("[PROGRESS] Per-node progress Excel files disabled")

    print(
        f"Web scrape mode: Selenium Grid "
        f"({grid_url}, auto_nodes={SELENIUM_AUTO_MANAGE_NODES}, "
        f"{MAX_SELENIUM_GRID_NODES} nodes x {SE_NODE_MAX_SESSIONS} sessions "
        f"= {SELENIUM_MAX_PARALLEL_SESSIONS} parallel)"
    )
    return _run_grid_scrape(
        df_input,
        grid_url,
        progress_dir=progress_dir,
        write_progress_files=write_progress_files,
        progress_callback=progress_callback,
    )


# Friendly aliases for imports from other scripts.
scrape_registration_date_and_fuel = scrape_kerala_reg_date_fuel
scrape_reg_date_fuel_type = scrape_kerala_reg_date_fuel


def _status_timestamp():
    return datetime.now().isoformat(timespec="seconds")


def load_folder_status(status_path):
    """Load or initialize the folder-level JSON status file."""
    status_path = Path(status_path)
    if not status_path.exists():
        return {
            "version": 1,
            "updated_at": _status_timestamp(),
            "files": {},
        }

    with status_path.open("r", encoding="utf-8") as fh:
        status = json.load(fh)

    if not isinstance(status, dict):
        raise ValueError(f"Invalid status JSON shape in {status_path}")
    status.setdefault("version", 1)
    status.setdefault("updated_at", _status_timestamp())
    status.setdefault("files", {})
    if not isinstance(status["files"], dict):
        raise ValueError(f"Invalid 'files' object in {status_path}")
    return status


def save_folder_status(status_path, status):
    """Atomically save the folder-level JSON status file."""
    status_path = Path(status_path)
    status_path.parent.mkdir(parents=True, exist_ok=True)
    status["updated_at"] = _status_timestamp()

    tmp_path = status_path.with_name(status_path.name + ".tmp")
    with tmp_path.open("w", encoding="utf-8") as fh:
        json.dump(status, fh, indent=2, sort_keys=True)
        fh.write("\n")
    tmp_path.replace(status_path)


def discover_workbooks(folder_path, recursive=False):
    """Find supported workbook files in a folder."""
    folder_path = Path(folder_path)
    iterator = folder_path.rglob("*") if recursive else folder_path.iterdir()
    workbooks = []
    for path in iterator:
        if not path.is_file():
            continue
        name = path.name
        # Skip Excel lock files and live-save temp workbooks.
        if name.startswith("~$") or name.startswith("."):
            continue
        if ".live_save.tmp" in name.lower() or ".final_save.tmp" in name.lower():
            continue
        if path.suffix.lower() in SUPPORTED_WORKBOOK_EXTENSIONS:
            workbooks.append(path)
    return sorted(workbooks, key=lambda p: str(p).lower())


def workbook_status_key(folder_path, workbook_path):
    """Use a stable relative path key inside the status JSON."""
    return (
        Path(workbook_path)
        .resolve()
        .relative_to(Path(folder_path).resolve())
        .as_posix()
    )


def save_workbook_frames_atomic(excel_path, sheets_dict):
    """Save workbook data through a temp file before replacing the source."""
    excel_path = Path(excel_path)
    tmp_path = excel_path.with_name(
        f".{excel_path.stem}.final_save.tmp{excel_path.suffix}"
    )
    with pd.ExcelWriter(tmp_path, engine="openpyxl") as writer:
        for sheet, data in sheets_dict.items():
            data.to_excel(writer, sheet_name=sheet, index=False)
    tmp_path.replace(excel_path)


def process_workbook_in_place(
    excel_path,
    write_progress_files=True,
    live_save=True,
    live_save_every=DEFAULT_LIVE_SAVE_EVERY,
):
    """Scrape all eligible sheets in a workbook and save back to the same file."""
    excel_path = Path(excel_path)
    if excel_path.suffix.lower() not in SUPPORTED_WORKBOOK_EXTENSIONS:
        raise ValueError(
            f"Unsupported workbook type '{excel_path.suffix}'. "
            f"Supported: {', '.join(sorted(SUPPORTED_WORKBOOK_EXTENSIONS))}"
        )
    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    progress_dir = make_progress_dir(excel_path) if write_progress_files else None

    with pd.ExcelFile(excel_path) as xls:
        sheets_dict = {
            sheet: pd.read_excel(xls, sheet_name=sheet, dtype=str)
            for sheet in xls.sheet_names
        }

    scraped_sheets = []
    skipped_sheets = []

    for sheet_name, df_sheet in sheets_dict.items():
        veh_col = find_vehicle_reg_column(df_sheet)
        if veh_col is None:
            print(
                f"[SKIP] Sheet '{sheet_name}': no vehicle registration header keyword found "
                f"(columns: {list(df_sheet.columns)})"
            )
            skipped_sheets.append(sheet_name)
            continue

        # Preserve already-filled detail columns; scrape only pending rows.
        for col in DETAIL_COLUMNS:
            if col not in df_sheet.columns:
                df_sheet[col] = ""
            else:
                df_sheet[col] = df_sheet[col].fillna("").astype(str).replace({"nan": ""})

        print(
            f"[OK] Sheet '{sheet_name}': found vehicle column '{veh_col}' "
            f"({len(df_sheet)} rows) - starting web scrape "
            f"(existing Registration Date / Fuel values are kept)"
        )
        live_updater = (
            ExcelLiveUpdater(
                excel_path,
                sheet_name,
                detail_columns=DETAIL_COLUMNS,
                save_every=live_save_every,
            )
            if live_save
            else None
        )
        try:
            sheets_dict[sheet_name] = scrape_kerala_reg_date_fuel(
                df_sheet,
                use_selenium_grid=USE_SELENIUM_GRID,
                progress_dir=progress_dir,
                write_progress_files=write_progress_files,
                progress_callback=live_updater.update_row if live_updater else None,
            )
        finally:
            if live_updater:
                live_updater.flush()
        scraped_sheets.append(sheet_name)

    if not scraped_sheets:
        print(
            "[ERROR] No sheet contained a vehicle registration header keyword "
            "(looked for columns with 'veh'+'reg' or 'vehicle'). Nothing scraped."
        )
        return {
            "saved": False,
            "scraped_sheets": scraped_sheets,
            "skipped_sheets": skipped_sheets,
            "progress_dir": progress_dir,
        }

    save_workbook_frames_atomic(excel_path, sheets_dict)

    print(f"[OK] Results saved back to '{excel_path}'")
    if progress_dir:
        print(f"[OK] Per-node progress files kept in: {progress_dir}")

    return {
        "saved": True,
        "scraped_sheets": scraped_sheets,
        "skipped_sheets": skipped_sheets,
        "progress_dir": progress_dir,
    }


def _is_file_complete(file_entry):
    return str(file_entry.get("status", "")).lower() in {"processed", "skipped"}


def _file_status(file_entry):
    return str((file_entry or {}).get("status", "pending") or "pending").lower()


def count_pending_vehicles_in_workbook(excel_path):
    """
    Count eligible vehicle rows that still need Registration Date / Fuel.

    Returns (pending_count, sheet_summaries). sheet_summaries lists
    {sheet, pending, total_eligible} for sheets that have a vehicle column.
    """
    excel_path = Path(excel_path)
    pending_total = 0
    sheet_summaries = []

    with pd.ExcelFile(excel_path) as xls:
        for sheet_name in xls.sheet_names:
            df = pd.read_excel(xls, sheet_name=sheet_name, dtype=str)
            veh_col = find_vehicle_reg_column(df)
            if veh_col is None:
                continue

            for col in DETAIL_COLUMNS:
                if col not in df.columns:
                    df[col] = ""

            df[veh_col] = df[veh_col].apply(normalize_vehicle_number)
            valid_mask = df[veh_col].apply(is_vehicle_number_eligible)
            completed_mask = df.apply(_row_has_completed_details, axis=1)
            pending = int((valid_mask & ~completed_mask).sum())
            eligible = int(valid_mask.sum())
            pending_total += pending
            sheet_summaries.append(
                {
                    "sheet": sheet_name,
                    "pending": pending,
                    "total_eligible": eligible,
                }
            )

    return pending_total, sheet_summaries


def sync_folder_status_from_disk(workbooks, status, folder_path):
    """
    Ensure every on-disk workbook has a JSON entry, then set status from data:

    - pending   : at least one eligible vehicle is missing Registration Date / Fuel
    - processed : all eligible vehicles already have data (existing values kept)
    - skipped   : workbook has no vehicle-registration sheet
    - failed    : workbook could not be read during the startup scan

    Returns list of (workbook, key, pending_count) for files that still need scraping.
    """
    now = _status_timestamp()
    created = 0
    marked_pending = 0
    marked_processed = 0
    marked_skipped = 0
    marked_failed = 0
    queue = []

    print("=" * 80)
    print("STARTUP SCAN: syncing folder files into status JSON")
    print("=" * 80)

    known_keys = set()
    for workbook in workbooks:
        key = workbook_status_key(folder_path, workbook)
        known_keys.add(key)
        file_entry = status["files"].get(key)
        if file_entry is None:
            file_entry = {
                "status": "pending",
                "path": key,
                "created_at": now,
            }
            status["files"][key] = file_entry
            created += 1
            print(f"[JSON] Created status entry: {key}")
        else:
            file_entry.setdefault("path", key)
            file_entry.setdefault("created_at", now)

        try:
            pending_count, sheet_summaries = count_pending_vehicles_in_workbook(workbook)
        except Exception as exc:
            file_entry.update(
                {
                    "status": "failed",
                    "path": key,
                    "updated_at": now,
                    "error": f"Startup scan failed: {exc}",
                    "pending_vehicles": None,
                }
            )
            marked_failed += 1
            print(f"[SCAN][FAIL] {key}: {exc}")
            continue

        eligible_total = sum(item["total_eligible"] for item in sheet_summaries)
        file_entry["pending_vehicles"] = pending_count
        file_entry["eligible_vehicles"] = eligible_total
        file_entry["path"] = key
        file_entry["updated_at"] = now
        file_entry["scan_at"] = now

        if not sheet_summaries:
            file_entry.update(
                {
                    "status": "skipped",
                    "error": "No vehicle registration sheet found",
                    "completed_at": now,
                }
            )
            marked_skipped += 1
            print(f"[SCAN] skipped (no vehicle sheet): {key}")
            continue

        if pending_count > 0:
            file_entry.update(
                {
                    "status": "pending",
                    "error": "",
                }
            )
            # Drop stale completion markers so the file is eligible again.
            file_entry.pop("completed_at", None)
            file_entry.pop("resume_note", None)
            marked_pending += 1
            queue.append((workbook, key, pending_count))
            print(
                f"[SCAN] pending: {key} "
                f"({pending_count} vehicle(s) still need data; existing values kept)"
            )
        else:
            file_entry.update(
                {
                    "status": "processed",
                    "saved": True,
                    "error": "",
                    "completed_at": now,
                    "resume_note": "Startup scan: all eligible vehicles already have data",
                }
            )
            marked_processed += 1
            print(
                f"[SCAN] processed: {key} "
                f"(0 pending; {eligible_total} vehicle(s) already filled)"
            )

    # Keep orphan JSON keys, but note files that disappeared from disk.
    for key, file_entry in status["files"].items():
        if key in known_keys:
            continue
        if str(file_entry.get("missing_on_disk", "")).lower() == "true":
            continue
        file_entry["missing_on_disk"] = True
        file_entry["updated_at"] = now
        print(f"[JSON] Entry exists but file missing on disk (left unchanged): {key}")

    print("-" * 80)
    print(f"Scan complete: created={created}, pending={marked_pending}, "
          f"processed={marked_processed}, skipped={marked_skipped}, failed={marked_failed}")
    print(f"Files queued for scraping: {len(queue)}")
    print("=" * 80)

    return queue


def process_folder(
    folder_path,
    status_path=None,
    recursive=False,
    write_progress_files=False,
    live_save=True,
    live_save_every=DEFAULT_LIVE_SAVE_EVERY,
):
    """
    Process every supported workbook in a folder, using JSON status to resume.

    Startup:
      1. Discover workbooks on disk.
      2. Create missing JSON entries.
      3. Mark each file pending/processed/skipped based on whether any
         eligible vehicle is still missing Registration Date / Fuel.
      4. Scrape queued files one after another.
         Already-filled vehicle rows are never cleared or re-scraped.
    """
    folder_path = Path(folder_path).resolve()
    if not folder_path.is_dir():
        raise NotADirectoryError(f"Folder not found: {folder_path}")

    if status_path is None:
        status_path = folder_path / DEFAULT_STATUS_FILENAME
    else:
        status_path = Path(status_path)
        if not status_path.is_absolute():
            status_path = folder_path / status_path

    workbooks = discover_workbooks(folder_path, recursive=recursive)
    status = load_folder_status(status_path)
    status["folder"] = str(folder_path)

    print("=" * 80)
    print("FOLDER PROCESSING")
    print("=" * 80)
    print(f"Folder: {folder_path}")
    print(f"Status JSON: {status_path}")
    print(f"Supported files found: {len(workbooks)}")
    print("=" * 80)

    queue = sync_folder_status_from_disk(workbooks, status, folder_path)
    save_folder_status(status_path, status)

    processed = 0
    skipped = 0
    failed = 0

    # Also retry files left as failed from an earlier run if they still exist.
    queued_keys = {key for _, key, _ in queue}
    for workbook in workbooks:
        key = workbook_status_key(folder_path, workbook)
        if key in queued_keys:
            continue
        if _file_status(status["files"].get(key, {})) == "failed":
            try:
                pending_count, _ = count_pending_vehicles_in_workbook(workbook)
            except Exception:
                pending_count = -1
            if pending_count != 0:
                queue.append((workbook, key, pending_count))
                queued_keys.add(key)

    # Stable one-file-after-another order.
    queue.sort(key=lambda item: str(item[1]).lower())

    for workbook, key, pending_hint in queue:
        file_entry = status["files"].setdefault(key, {"path": key})

        try:
            pending_count, _ = count_pending_vehicles_in_workbook(workbook)
        except Exception as exc:
            print(
                f"[WARN] Could not re-count pending vehicles in {key}: {exc}. "
                "Will scrape remaining rows only.",
                flush=True,
            )
            pending_count = pending_hint if pending_hint else -1

        if pending_count == 0:
            file_entry.update(
                {
                    "status": "processed",
                    "path": key,
                    "completed_at": _status_timestamp(),
                    "updated_at": _status_timestamp(),
                    "saved": True,
                    "error": "",
                    "pending_vehicles": 0,
                    "resume_note": "No pending vehicles; existing data kept",
                }
            )
            save_folder_status(status_path, status)
            processed += 1
            print(f"[SKIP] No pending vehicles (existing data kept): {key}")
            continue

        print("\n" + "=" * 80)
        print(f"[FILE] Processing: {key}")
        print(f"[RESUME] Pending vehicles only: {pending_count}")
        print("Existing Registration Date / Fuel values will be preserved.")
        print("=" * 80)

        file_entry.update(
            {
                "status": "processing",
                "path": key,
                "started_at": _status_timestamp(),
                "updated_at": _status_timestamp(),
                "pending_vehicles": pending_count,
                "error": "",
            }
        )
        save_folder_status(status_path, status)

        try:
            result = process_workbook_in_place(
                workbook,
                write_progress_files=write_progress_files,
                live_save=live_save,
                live_save_every=live_save_every,
            )
            # Re-check after scrape so partial failures stay pending/failed correctly.
            try:
                remaining_after, _ = count_pending_vehicles_in_workbook(workbook)
            except Exception:
                remaining_after = 0 if result["saved"] else -1

            if result["saved"] and remaining_after == 0:
                next_status = "processed"
            elif result["saved"] and remaining_after > 0:
                next_status = "pending"
            elif not result["saved"]:
                next_status = "skipped"
            else:
                next_status = "failed"

            file_entry.update(
                {
                    "status": next_status,
                    "completed_at": _status_timestamp(),
                    "updated_at": _status_timestamp(),
                    "saved": result["saved"],
                    "scraped_sheets": result["scraped_sheets"],
                    "skipped_sheets": result["skipped_sheets"],
                    "progress_dir": result["progress_dir"],
                    "pending_vehicles": (
                        remaining_after if remaining_after >= 0 else None
                    ),
                    "error": (
                        ""
                        if next_status in {"processed", "pending"}
                        else (
                            "No vehicle registration sheet found"
                            if next_status == "skipped"
                            else "Scrape finished with remaining/unknown pending rows"
                        )
                    ),
                }
            )
            if next_status == "processed":
                processed += 1
            elif next_status == "skipped":
                skipped += 1
            elif next_status == "pending":
                # Partial progress kept; file stays eligible next run.
                print(
                    f"[PARTIAL] {key}: {remaining_after} vehicle(s) still pending; "
                    "existing filled data kept"
                )
            else:
                failed += 1
            save_folder_status(status_path, status)
        except Exception as exc:
            failed += 1
            file_entry.update(
                {
                    "status": "failed",
                    "completed_at": _status_timestamp(),
                    "updated_at": _status_timestamp(),
                    "error": str(exc),
                }
            )
            save_folder_status(status_path, status)
            print(f"[ERROR] Failed processing {key}: {exc}")

    already_complete = sum(
        1
        for file_entry in status["files"].values()
        if _is_file_complete(file_entry)
    )

    print("\n" + "=" * 80)
    print("FOLDER PROCESSING COMPLETE")
    print("=" * 80)
    print(f"Processed now: {processed}")
    print(f"Already complete / skipped in scan: {already_complete}")
    print(f"Failed: {failed}")
    print(f"Status JSON: {status_path}")
    print("=" * 80)

    return {
        "processed": processed,
        "skipped": skipped,
        "failed": failed,
        "status_path": str(status_path),
        "queued": len(queue),
        "already_complete": already_complete,
    }


def parse_args(argv=None):
    parser = argparse.ArgumentParser(
        description=(
            "Scrape Kerala Registration Date and Fuel into Excel files. "
            "Set REG_DATE_FUEL_EXCEL_PATH in .env for normal use. "
            "Command-line args remain optional overrides."
        )
    )
    parser.add_argument(
        "path",
        nargs="?",
        default=INPUT_PATH or None,
        help=(
            "Optional override for REG_DATE_FUEL_EXCEL_PATH "
            "(.xlsx file or folder)."
        ),
    )
    parser.add_argument(
        "--status-file",
        default=STATUS_FILE,
        help=(
            "Optional override for STATUS_FILE. Relative paths are placed "
            "inside the input folder."
        ),
    )
    parser.add_argument(
        "--recursive",
        action="store_true",
        default=RECURSIVE,
        help="Optional override: folder mode only, scan subfolders too.",
    )
    parser.add_argument(
        "--progress-files",
        choices=("auto", "yes", "no"),
        default=PROGRESS_FILES,
        help="Optional override for PROGRESS_FILES: auto, yes, or no.",
    )
    parser.add_argument(
        "--live-save",
        choices=("yes", "no"),
        default="yes" if LIVE_SAVE else "no",
        help="Optional override for LIVE_SAVE: yes or no.",
    )
    parser.add_argument(
        "--live-save-every",
        type=int,
        default=LIVE_SAVE_EVERY,
        help="Optional override for LIVE_SAVE_EVERY.",
    )
    return parser.parse_args(argv)


def main(argv=None):
    args = parse_args(argv)
    if not args.path:
        raise FileNotFoundError(
            "Input path not set. Set REG_DATE_FUEL_EXCEL_PATH in .env "
            "or pass a workbook/folder path."
        )

    input_path = Path(args.path)
    if not input_path.exists():
        raise FileNotFoundError(
            f"Input path not found: {input_path}. Pass a workbook/folder path "
            "or set REG_DATE_FUEL_EXCEL_PATH in .env."
        )

    if args.progress_files == "yes":
        write_progress_files = True
    elif args.progress_files == "no":
        write_progress_files = False
    else:
        write_progress_files = input_path.is_file()

    live_save = args.live_save == "yes"
    live_save_every = max(1, int(args.live_save_every or 1))

    if input_path.is_dir():
        process_folder(
            input_path,
            status_path=args.status_file,
            recursive=args.recursive,
            write_progress_files=write_progress_files,
            live_save=live_save,
            live_save_every=live_save_every,
        )
    else:
        process_workbook_in_place(
            input_path,
            write_progress_files=write_progress_files,
            live_save=live_save,
            live_save_every=live_save_every,
        )


if __name__ == "__main__":
    main()
